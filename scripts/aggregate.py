# -*- coding: utf-8 -*-
"""
aggregate.py —— 班级聚合：挑学生佳句、注入场景词汇、生成班级成绩台账。

用法：
    python aggregate.py --results-dir ./results --assignment assignment.json \
                        --outdir ./final --ledger 班级台账.xlsx

职责：
  1. 读取每名学生批改 JSON（results/*.json）
  2. **按班级**分组，各自按得分降序挑佳句，注入该班每份 JSON 的 class_shared.quotes
  3. 注入场景词汇（assignment.json 的 scene_vocab；若有 scene_vocab_by_class 则按班覆盖）
  4. 输出回填后的 JSON 到 outdir，供 render.py 渲染
  5. 生成班级成绩台账 Excel（含得分分布、错误类型统计、各班均分）

重要：学生佳句与场景词汇都是**班级级**内容——同一个班的学生共享**同一组**佳句，
换一个班，佳句的人名、句序以及场景词汇的措辞都会不同。
因此绝不能跨班聚合，否则会把别班学生的句子塞进本班报告。

关于「学生自己的句子」：佳句区**允许出现学生本人的句子**。
它取的是本班所有学生里得分最高的前 N 句，学生自己写得好，就会出现在自己的报告里。
同一班内每份报告的佳句列表**完全一致**（这是班级级内容的定义），不要再按人做过滤。
"""
import argparse
import collections
import json
import os


def load_all(d):
    out = []
    for fn in sorted(os.listdir(d)):
        if fn.endswith(".json"):
            with open(os.path.join(d, fn), "r", encoding="utf-8") as f:
                out.append((os.path.join(d, fn), fn, json.load(f)))
    return out


def rank_quotes(students, limit):
    """把本班佳句按得分降序排成一个池子：每生至多 1 句，取前 limit 条。

    池子是**班级级**的——本班每份报告都用同一份列表，其中可以包含学生自己的句子。
    """
    cands = []
    for _, _, s in students:
        hl = s.get("highlights", {}).get("items", [])
        if not hl:
            continue
        sent = (hl[0].get("sentence") or "").strip()
        if len(sent) < 15:
            continue
        name = s.get("student", {}).get("name", "")
        if name:
            cands.append((s.get("score") or 0, name, sent))
    cands.sort(key=lambda x: (-x[0], x[1]))
    seen, out = set(), []
    for _, name, sent in cands:
        if name in seen:
            continue
        seen.add(name)
        out.append({"name": name, "sentence": sent})
        if len(out) >= limit:
            break
    return out


def err_stats(students):
    c = collections.Counter()
    for _, _, s in students:
        for e in s.get("errors", []):
            for n in e.get("notes", []):
                c[n.get("type") or "未分类"] += 1
    return c


def score_bands(students):
    bands = collections.Counter()
    for _, _, s in students:
        v = s.get("score")
        if v is None:
            bands["缺"] += 1
        elif v >= 20:
            bands["20-25"] += 1
        elif v >= 15:
            bands["15-19"] += 1
        elif v >= 10:
            bands["10-14"] += 1
        else:
            bands["0-9"] += 1
    return bands


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results-dir", required=True)
    ap.add_argument("--assignment", required=True, help="题目级配置 assignment.json")
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--ledger", help="输出班级成绩台账 xlsx 路径")
    ap.add_argument("--quote-count", type=int, default=6, help="每班挑几条佳句")
    args = ap.parse_args()

    with open(args.assignment, "r", encoding="utf-8") as f:
        asg = json.load(f)

    students = load_all(args.results_dir)

    # --- 按班级分组 ---
    by_class = collections.defaultdict(list)
    for p, fn, s in students:
        cls = (s.get("student", {}).get("class") or "未分班").strip()
        by_class[cls].append((p, fn, s))

    class_pool, class_vocab = {}, {}
    for cls, group in by_class.items():
        class_pool[cls] = rank_quotes(group, args.quote_count)
        class_vocab[cls] = (asg.get("scene_vocab_by_class", {}) or {}).get(cls) or asg.get("scene_vocab", [])
        print("  %s：%d 人，佳句池 %d 条，场景 %d 组" % (cls, len(group), len(class_pool[cls]), len(class_vocab[cls])))

    os.makedirs(args.outdir, exist_ok=True)
    for p, fn, s in students:
        s.setdefault("meta", {})
        for k in ("assignment", "title_date", "date", "weekday"):
            if asg.get(k):
                s["meta"].setdefault(k, asg[k])
        cls = (s.get("student", {}).get("class") or "未分班").strip()
        # 班级级：本班所有人共用同一份佳句列表（可包含学生自己的句子）
        s["class_shared"] = {
            "quotes": class_pool[cls],
            "vocab": class_vocab[cls],
        }
        with open(os.path.join(args.outdir, fn), "w", encoding="utf-8") as f:
            json.dump(s, f, ensure_ascii=False, indent=2)

    print("已回填 %d 份，覆盖 %d 个班" % (len(students), len(by_class)))

    if args.ledger:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:
            print("未安装 openpyxl，跳过台账。pip install openpyxl")
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "成绩台账"
        ws.append(["学号", "姓名", "班级", "得分", "词数", "语法错误数", "卷面等级", "问题标签", "主要错误类型"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for _, _, s in sorted(students, key=lambda x: (x[2].get("student", {}).get("class", ""),
                                                       x[2].get("student", {}).get("id", ""))):
            st = s.get("student", {})
            ws.append([
                st.get("id", ""), st.get("name", ""), st.get("class", ""),
                s.get("score", ""), s.get("word_count", ""), len(s.get("errors", [])),
                s.get("face", {}).get("grade", ""), "；".join(s.get("tags", [])),
                "；".join(sorted({n.get("type", "") for e in s.get("errors", [])
                                  for n in e.get("notes", []) if n.get("type")})),
            ])
        for col, w in zip("ABCDEFGHI", (10, 10, 8, 8, 8, 12, 10, 34, 40)):
            ws.column_dimensions[col].width = w

        ws2 = wb.create_sheet("班级诊断")
        ws2.append(["得分分布"])
        for k, v in sorted(score_bands(students).items()):
            ws2.append([k, v])
        ws2.append([])
        ws2.append(["错误类型", "人次"])
        for k, v in err_stats(students).most_common():
            ws2.append([k, v])
        ws2.append([])
        ws2.append(["班级", "人数", "均分"])
        for cls, group in sorted(by_class.items()):
            sc = [g[2].get("score") for g in group if isinstance(g[2].get("score"), (int, float))]
            ws2.append([cls, len(group), round(sum(sc) / len(sc), 1) if sc else ""])
        ws2.column_dimensions["A"].width = 24
        ws2.column_dimensions["B"].width = 10
        ws2.column_dimensions["C"].width = 10

        wb.save(args.ledger)
        print("台账已写入 %s" % args.ledger)


if __name__ == "__main__":
    main()
